"""Export visible columns from the Excel catalogue for the static website.

Run this script after replacing/updating Bibliotheek_Lijst_Michel.xlsx.
Columns hidden in Excel are omitted from both the output and the website.
"""

from __future__ import annotations

import json
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "Bibliotheek_Lijst_Michel.xlsx"
JSON_DESTINATION = ROOT / "data" / "books.json"
SCRIPT_DESTINATION = ROOT / "data" / "books.js"


def clean(value):
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return int(value)
    if isinstance(value, str):
        return " ".join(value.split())
    return value


def main() -> None:
    workbook = load_workbook(SOURCE, data_only=True, read_only=False)
    sheet = workbook.active
    visible_columns = [
        column
        for column in range(1, sheet.max_column + 1)
        if not sheet.column_dimensions[sheet.cell(1, column).column_letter].hidden
    ]
    headers = [clean(sheet.cell(1, column).value) for column in visible_columns]

    rows = []
    for row_number in range(2, sheet.max_row + 1):
        rows.append([clean(sheet.cell(row_number, column).value) for column in visible_columns])

    payload = {
        "sheet": sheet.title,
        "headers": headers,
        "rows": len(rows),
        "hiddenColumnsExcluded": sheet.max_column - len(visible_columns),
        "books": rows,
    }
    JSON_DESTINATION.parent.mkdir(exist_ok=True)
    serialized = json.dumps(payload, ensure_ascii=False, separators=(",", ":"))
    JSON_DESTINATION.write_text(serialized, encoding="utf-8")
    # A classic script is more robust than fetch() when the site is embedded in
    # a sandboxed third-party iframe. Escape the two JavaScript line separators.
    script_data = serialized.replace("\u2028", "\\u2028").replace("\u2029", "\\u2029")
    SCRIPT_DESTINATION.write_text(
        f"globalThis.BOOK_CATALOG={script_data};\n",
        encoding="utf-8",
    )
    print(
        f"Exported {len(rows):,} rows with {len(visible_columns)} visible columns; "
        f"excluded {sheet.max_column - len(visible_columns)} hidden columns."
    )


if __name__ == "__main__":
    main()
